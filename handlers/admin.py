from html import escape

from aiogram import F, Router
from aiogram.fsm.context import FSMContext
from aiogram.types import CallbackQuery, FSInputFile, Message

import database as db
import keyboards as kb
from states import AddAdminStates, AddVideoStates, AdjustVotesStates, BroadcastStates, EditVideoStates, VotingPeriodStates
from utils import DATE_FORMAT, is_admin, is_super_admin, parse_datetime

router = Router()


# ==================== ADMIN MENU ====================

@router.message(F.text == "🔐 Admin panel")
async def admin_panel(message: Message):
    if not await is_admin(message.from_user.id):
        return
    await message.answer(
        "🔐 <b>ADMIN PANEL</b>\n\nXush kelibsiz.",
        reply_markup=kb.admin_menu_keyboard(is_super_admin(message.from_user.id)),
    )


@router.message(F.text == "⬅️ Asosiy menyu")
async def back_to_main_menu(message: Message, state: FSMContext):
    await state.clear()
    await message.answer(
        "Asosiy menyu.", reply_markup=kb.main_menu_keyboard(await is_admin(message.from_user.id))
    )


# ==================== VIDEO QO'SHISH ====================

@router.message(F.text == "➕ Video qo'shish")
async def add_video_start(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    await state.set_state(AddVideoStates.waiting_for_video)
    await message.answer("🎬 Videoni yuboring.")


@router.message(AddVideoStates.waiting_for_video, F.video)
async def add_video_receive_video(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    await state.update_data(file_id=message.video.file_id)
    await state.set_state(AddVideoStates.waiting_for_title)
    await message.answer("📝 Ishtirokchi nomini yuboring.")


@router.message(AddVideoStates.waiting_for_video)
async def add_video_wrong_type(message: Message):
    if not await is_admin(message.from_user.id):
        return
    await message.answer("⚠️ Iltimos, video fayl yuboring.")


@router.message(AddVideoStates.waiting_for_title, F.text)
async def add_video_receive_title(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    await state.update_data(title=message.text)
    await state.set_state(AddVideoStates.waiting_for_description)
    await message.answer(kb.SKIP_DESCRIPTION_HINT)


@router.message(AddVideoStates.waiting_for_description, F.text)
async def add_video_receive_description(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    await state.update_data(description=message.text)
    await state.set_state(AddVideoStates.waiting_for_position)
    await message.answer("🔢 Tartib raqamini kiriting.\n\nMasalan: 1")


@router.message(AddVideoStates.waiting_for_position, F.text)
async def add_video_receive_position(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    try:
        position = int(message.text.strip())
    except ValueError:
        await message.answer("⚠️ Iltimos, tartib raqamini butun son sifatida kiriting. Masalan: 1")
        return

    data = await state.get_data()
    await db.add_video(data["file_id"], data["title"], data.get("description"), position)
    await state.clear()
    await message.answer(
        "✅ Video muvaffaqiyatli qo'shildi.",
        reply_markup=kb.admin_menu_keyboard(is_super_admin(message.from_user.id)),
    )


# ==================== VIDEOLAR RO'YXATI (ADMIN) ====================

@router.message(F.text == "📋 Videolar")
async def admin_list_videos(message: Message):
    if not await is_admin(message.from_user.id):
        return
    videos = await db.get_all_videos_including_inactive()
    if not videos:
        await message.answer("😔 Hozircha videolar mavjud emas.")
        return
    for video in videos:
        title = escape(video["title"])
        description = escape(video["description"] or "")
        status = "✅ Faol" if video["is_active"] else "🚫 Yashiringan"
        caption = (
            f"🎬 <b>{title}</b> ({video['position']}-o'rin) — {status}\n"
            f"{description}\n\n"
            f"🗳 {video['votes_count']} ta ovoz"
        ).strip()
        reply_markup = None
        if not video["is_active"]:
            reply_markup = kb.restore_video_keyboard(video["id"])
        await message.answer_video(video=video["file_id"], caption=caption, reply_markup=reply_markup)


@router.callback_query(F.data.startswith("restore_video:"))
async def cb_restore_video(callback: CallbackQuery):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    video_id = int(callback.data.split(":")[1])
    await db.restore_video(video_id)
    await callback.answer("✅ Video va uning ovozlari tiklandi!", show_alert=True)


# ==================== STATISTIKA ====================

@router.message(F.text == "📊 Statistika")
async def admin_stats(message: Message):
    if not await is_admin(message.from_user.id):
        return
    stats = await db.get_stats()
    await message.answer(
        "📊 <b>Statistika</b>\n\n"
        f"👥 Jami foydalanuvchilar: <b>{stats['users']}</b>\n"
        f"🎬 Jami videolar: <b>{stats['videos']}</b>\n"
        f"🗳 Jami berilgan ovozlar: <b>{stats['votes']}</b>\n"
        f"✅ Ovoz bergan foydalanuvchilar: <b>{stats['voted_users']}</b>\n"
        f"⏳ Hali ovoz bermagan foydalanuvchilar: <b>{stats['not_voted_users']}</b>"
    )


# ==================== VIDEO O'CHIRISH ====================

@router.message(F.text == "🗑 Video o'chirish")
async def admin_delete_video_list(message: Message):
    if not await is_admin(message.from_user.id):
        return
    videos = await db.get_all_videos()
    if not videos:
        await message.answer("😔 O'chirish uchun video mavjud emas.")
        return
    await message.answer(
        "O'chirmoqchi bo'lgan videongizni tanlang:",
        reply_markup=kb.delete_video_list_keyboard(videos),
    )


@router.callback_query(F.data.startswith("delete_ask:"))
async def cb_delete_ask(callback: CallbackQuery):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    video_id = int(callback.data.split(":")[1])
    video = await db.get_video(video_id)
    if not video:
        await callback.answer("⚠️ Video topilmadi.", show_alert=True)
        return
    await callback.message.answer(
        f"⚠️ Ushbu videoni yashirmoqchimisiz?\n\n🎬 <b>{escape(video['title'])}</b>\n\n"
        "(Ovozlari saqlanib qoladi, kerak bo'lsa keyinroq tiklashingiz mumkin.)",
        reply_markup=kb.confirm_delete_keyboard(video_id),
    )
    await callback.answer()


@router.callback_query(F.data.startswith("delete_confirm:"))
async def cb_delete_confirm(callback: CallbackQuery):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    video_id = int(callback.data.split(":")[1])
    await db.delete_video(video_id)
    await callback.message.edit_text("🚫 Video yashirildi (ovozlari saqlanib qoldi, \"📋 Videolar\"dan tiklash mumkin).")
    await callback.answer()


@router.callback_query(F.data == "delete_cancel")
async def cb_delete_cancel(callback: CallbackQuery):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    await callback.message.edit_text("❌ Bekor qilindi.")
    await callback.answer()


# ==================== VIDEO TAHRIRLASH ====================

@router.message(F.text == "✏️ Video tahrirlash")
async def admin_edit_video_list(message: Message):
    if not await is_admin(message.from_user.id):
        return
    videos = await db.get_all_videos()
    if not videos:
        await message.answer("😔 Tahrirlash uchun video mavjud emas.")
        return
    await message.answer(
        "Tahrirlamoqchi bo'lgan videongizni tanlang:",
        reply_markup=kb.edit_video_list_keyboard(videos),
    )


@router.callback_query(F.data.startswith("edit_pick:"))
async def cb_edit_pick(callback: CallbackQuery):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    video_id = int(callback.data.split(":")[1])
    video = await db.get_video(video_id)
    if not video:
        await callback.answer("⚠️ Video topilmadi.", show_alert=True)
        return
    await callback.message.answer(
        f"🎬 <b>{escape(video['title'])}</b>\n\nNimani tahrirlamoqchisiz?",
        reply_markup=kb.edit_field_keyboard(video_id),
    )
    await callback.answer()


@router.callback_query(F.data.startswith("edit_field:"))
async def cb_edit_field(callback: CallbackQuery, state: FSMContext):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    _, video_id_raw, field = callback.data.split(":")
    await state.update_data(edit_video_id=int(video_id_raw), edit_field=field)
    await state.set_state(EditVideoStates.waiting_for_new_value)

    prompts = {
        "title": "📝 Yangi nomni yuboring.",
        "description": "📄 Yangi tavsifni yuboring.",
        "position": "🔢 Yangi tartib raqamini yuboring.",
    }
    await callback.message.answer(prompts[field])
    await callback.answer()


@router.callback_query(F.data == "edit_cancel")
async def cb_edit_cancel(callback: CallbackQuery, state: FSMContext):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    await state.clear()
    await callback.message.edit_text("❌ Bekor qilindi.")
    await callback.answer()


@router.message(EditVideoStates.waiting_for_new_value, F.text)
async def edit_video_apply(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    data = await state.get_data()
    video_id: int = data["edit_video_id"]
    field: str = data["edit_field"]

    if field == "position":
        try:
            value: object = int(message.text.strip())
        except ValueError:
            await message.answer("⚠️ Iltimos, butun son kiriting.")
            return
        await db.update_video(video_id, position=value)
    elif field == "title":
        await db.update_video(video_id, title=message.text)
    elif field == "description":
        await db.update_video(video_id, description=message.text)

    await state.clear()
    await message.answer(
        "✅ Video muvaffaqiyatli yangilandi.",
        reply_markup=kb.admin_menu_keyboard(is_super_admin(message.from_user.id)),
    )


# ==================== OVOZ BERISH MUDDATI ====================

@router.message(F.text == "📅 Sana belgilash")
async def voting_period_start(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    start_iso, end_iso = await db.get_voting_period()
    current = ""
    if start_iso and end_iso:
        from datetime import datetime as dt
        current = (
            f"\n\nHozirgi muddat:\n"
            f"Boshlanish: {dt.fromisoformat(start_iso).strftime(DATE_FORMAT)}\n"
            f"Tugash: {dt.fromisoformat(end_iso).strftime(DATE_FORMAT)}"
        )
    await state.set_state(VotingPeriodStates.waiting_for_start)
    await message.answer(
        "📅 Ovoz berish <b>boshlanish</b> sanasini kiriting.\n\n"
        "Format: kun.oy.yil soat:daqiqa\n"
        "Masalan: 20.09.2026 09:00" + current
    )


@router.message(VotingPeriodStates.waiting_for_start, F.text)
async def voting_period_receive_start(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    start_dt = parse_datetime(message.text)
    if not start_dt:
        await message.answer("⚠️ Format xato. Masalan: 20.09.2026 09:00")
        return
    await state.update_data(start_iso=start_dt.isoformat())
    await state.set_state(VotingPeriodStates.waiting_for_end)
    await message.answer("📅 Endi ovoz berish <b>tugash</b> sanasini kiriting.\n\nMasalan: 25.09.2026 18:00")


@router.message(VotingPeriodStates.waiting_for_end, F.text)
async def voting_period_receive_end(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    end_dt = parse_datetime(message.text)
    if not end_dt:
        await message.answer("⚠️ Format xato. Masalan: 25.09.2026 18:00")
        return

    data = await state.get_data()
    from datetime import datetime as dt
    start_dt = dt.fromisoformat(data["start_iso"])

    if end_dt <= start_dt:
        await message.answer("⚠️ Tugash sanasi boshlanish sanasidan keyin bo'lishi kerak. Qaytadan kiriting.")
        return

    await db.set_voting_period(start_dt.isoformat(), end_dt.isoformat())
    await db.set_setting("reminder_sent_for", "")
    await state.clear()
    await message.answer(
        "✅ Ovoz berish muddati belgilandi.\n\n"
        f"Boshlanish: <b>{start_dt.strftime(DATE_FORMAT)}</b>\n"
        f"Tugash: <b>{end_dt.strftime(DATE_FORMAT)}</b>",
        reply_markup=kb.admin_menu_keyboard(is_super_admin(message.from_user.id)),
    )


# ==================== ADMINLARNI BOSHQARISH (faqat asosiy admin) ====================

@router.message(F.text == "👤 Adminlar")
async def admins_menu(message: Message):
    if not is_super_admin(message.from_user.id):
        return
    admin_ids = await db.get_admin_ids()
    lines = ["👤 <b>ADMINLAR</b>\n"]
    lines.append(f"🔑 Asosiy admin: <code>{message.from_user.id}</code> (siz)")
    if admin_ids:
        for admin_id in admin_ids:
            lines.append(f"👤 <code>{admin_id}</code>")
    else:
        lines.append("\nQo'shimcha adminlar hali yo'q.")
    await message.answer("\n".join(lines), reply_markup=kb.admins_list_keyboard(admin_ids))


@router.message(F.text == "➕ Admin qo'shish")
async def add_admin_start(message: Message, state: FSMContext):
    if not is_super_admin(message.from_user.id):
        return
    await state.set_state(AddAdminStates.waiting_for_id)
    await message.answer("🆔 Yangi adminning Telegram ID raqamini yuboring.")


@router.message(AddAdminStates.waiting_for_id, F.text)
async def add_admin_receive_id(message: Message, state: FSMContext):
    if not is_super_admin(message.from_user.id):
        return
    await state.clear()
    try:
        new_admin_id = int(message.text.strip())
    except ValueError:
        await message.answer("⚠️ Iltimos, to'g'ri Telegram ID (raqam) kiriting.")
        return

    added = await db.add_admin(new_admin_id, message.from_user.id)
    if added:
        await message.answer(f"✅ <code>{new_admin_id}</code> admin sifatida qo'shildi.")
    else:
        await message.answer("⚠️ Bu foydalanuvchi allaqachon admin.")


@router.callback_query(F.data.startswith("admin_remove:"))
async def cb_admin_remove(callback: CallbackQuery):
    if not is_super_admin(callback.from_user.id):
        await callback.answer()
        return
    admin_id = int(callback.data.split(":")[1])
    await db.remove_admin(admin_id)
    await callback.message.edit_text(f"🗑 <code>{admin_id}</code> adminlikdan olindi.")
    await callback.answer()


# ==================== BROADCAST (barchaga xabar yuborish) ====================

@router.message(F.text == "📢 Xabar yuborish")
async def broadcast_start(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    await state.set_state(BroadcastStates.waiting_for_message)
    await message.answer("📢 Barcha foydalanuvchilarga yuboriladigan xabar matnini kiriting.")


@router.message(BroadcastStates.waiting_for_message, F.text)
async def broadcast_receive_message(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    await state.update_data(broadcast_text=message.text)
    await message.answer(
        f"Quyidagi xabar barcha foydalanuvchilarga yuborilsinmi?\n\n{escape(message.text)}",
        reply_markup=kb.broadcast_confirm_keyboard(),
    )


@router.callback_query(F.data == "broadcast_send")
async def cb_broadcast_send(callback: CallbackQuery, state: FSMContext):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    data = await state.get_data()
    text = data.get("broadcast_text")
    await state.clear()

    if not text:
        await callback.answer()
        return

    await callback.message.edit_text("📤 Yuborilmoqda...")

    user_ids = await db.get_all_user_telegram_ids()
    sent, failed = 0, 0
    for telegram_id in user_ids:
        try:
            await callback.bot.send_message(telegram_id, text)
            sent += 1
        except Exception:
            failed += 1

    await callback.message.answer(
        f"✅ Yuborish yakunlandi.\n\n📤 Yuborildi: {sent}\n⚠️ Yetkazilmadi: {failed}"
    )
    await callback.answer()


@router.callback_query(F.data == "broadcast_cancel")
async def cb_broadcast_cancel(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await callback.message.edit_text("❌ Bekor qilindi.")
    await callback.answer()


# ==================== EXCEL EKSPORT ====================

@router.message(F.text == "📤 Excel eksport")
async def export_excel(message: Message):
    if not await is_admin(message.from_user.id):
        return

    await message.answer("📤 Fayl tayyorlanmoqda...")

    from openpyxl import Workbook

    wb = Workbook()

    results_sheet = wb.active
    results_sheet.title = "Natijalar"
    results_sheet.append(
        ["O'rin", "Ishtirokchi", "Organik ovozlar", "Tuzatish ovozlari", "Jami ovozlar", "Foiz"]
    )
    videos = await db.get_results()
    total_votes = sum(v["votes_count"] for v in videos)
    for i, video in enumerate(videos, start=1):
        percent = (video["votes_count"] / total_votes * 100) if total_votes > 0 else 0
        adjustment_total = video["votes_count"] - video["organic_votes_count"]
        results_sheet.append([
            i,
            video["title"],
            video["organic_votes_count"],
            adjustment_total,
            video["votes_count"],
            round(percent, 1),
        ])

    users_sheet = wb.create_sheet("Foydalanuvchilar")
    users_sheet.append(["Ism", "Familiya", "Username", "Telegram ID", "Ovoz bergan", "Ro'yxatdan o'tgan"])
    users = await db.get_all_users_full()
    for user in users:
        users_sheet.append([
            user.get("first_name") or "",
            user.get("last_name") or "",
            f"@{user['username']}" if user.get("username") else "",
            user["telegram_id"],
            user.get("voted_title") or "Hali ovoz bermagan",
            user["created_at"].split("T")[0],
        ])

    adjustments_sheet = wb.create_sheet("Ovoz tuzatishlari")
    adjustments_sheet.append(
        ["Sana", "Ishtirokchi", "Miqdor", "Admin ID", "Sabab", "Holati"]
    )
    for adj in await db.get_all_adjustments():
        adjustments_sheet.append([
            adj["created_at"].split("T")[0],
            adj["video_title"],
            adj["amount"],
            adj["admin_id"],
            adj["reason"],
            "Bekor qilingan" if adj["is_cancelled"] else "Faol",
        ])

    file_path = "/tmp/natijalar.xlsx"
    wb.save(file_path)

    await message.answer_document(
        FSInputFile(file_path, filename="natijalar.xlsx"),
        caption="📊 Natijalar va foydalanuvchilar ro'yxati.",
    )


# ==================== OVOZ TUZATISH (texnik nosozlik sabab yo'qolgan ovozlarni tiklash) ====================
# MUHIM: Bu funksiya ishtirokchiga asossiz ustunlik berish uchun EMAS.
# Faqat texnik xatolik sabab hisoblanmay qolgan ovozlarni, to'liq shaffof va
# hujjatlashtirilgan tarzda qayta hisobga olish uchun.

@router.message(F.text == "⚖️ Ovoz tuzatish")
async def adjust_votes_start(message: Message):
    if not await is_admin(message.from_user.id):
        return
    videos = await db.get_all_videos()
    if not videos:
        await message.answer("😔 Video mavjud emas.")
        return
    await message.answer(
        "⚖️ <b>Ovoz tuzatish</b>\n\n"
        "Bu funksiya FAQAT texnik nosozlik sabab yo'qolgan ovozlarni qayta "
        "hisobga olish uchun. Har bir tuzatish kim, qachon, nima sababdan "
        "qo'shganini ko'rsatib, tarixda saqlanadi.\n\n"
        "Qaysi ishtirokchiga tuzatish kiritmoqchisiz?",
        reply_markup=kb.adjust_video_pick_keyboard(videos),
    )


@router.callback_query(F.data.startswith("adjust_pick:"))
async def cb_adjust_pick(callback: CallbackQuery, state: FSMContext):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    video_id = int(callback.data.split(":")[1])
    video = await db.get_video(video_id)
    if not video:
        await callback.answer("⚠️ Video topilmadi.", show_alert=True)
        return
    await state.update_data(adjust_video_id=video_id)
    await state.set_state(AdjustVotesStates.waiting_for_amount)
    await callback.message.answer(
        f"🎬 <b>{escape(video['title'])}</b>\n"
        f"📊 Hozirgi jami ovozlar: <b>{video['votes_count']}</b>\n\n"
        "Nechta ovoz qo'shmoqchisiz? (butun son kiriting, masalan: 37)"
    )
    await callback.answer()


@router.message(AdjustVotesStates.waiting_for_amount, F.text)
async def adjust_votes_receive_amount(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    try:
        amount = int(message.text.strip())
    except ValueError:
        await message.answer("⚠️ Iltimos, butun son kiriting. Masalan: 37")
        return
    if amount == 0:
        await message.answer("⚠️ Son 0 dan farqli bo'lishi kerak.")
        return

    await state.update_data(adjust_amount=amount)
    await state.set_state(AdjustVotesStates.waiting_for_reason)
    await message.answer(
        "📝 Bu tuzatishning sababini yozing.\n\n"
        "Masalan: \"Texnik nosozlik sabab yo'qolgan ovozlarni qayta hisobga olish.\""
    )


@router.message(AdjustVotesStates.waiting_for_reason, F.text)
async def adjust_votes_receive_reason(message: Message, state: FSMContext):
    if not await is_admin(message.from_user.id):
        return
    data = await state.get_data()
    video_id = data["adjust_video_id"]
    amount = data["adjust_amount"]
    reason = message.text.strip()

    await state.update_data(adjust_reason=reason)

    video = await db.get_video(video_id)
    old_total = video["votes_count"]
    new_total = old_total + amount
    sign = "+" if amount > 0 else ""

    await message.answer(
        f"⚖️ <b>Tasdiqlashni kuting</b>\n\n"
        f"🎬 Ishtirokchi: <b>{escape(video['title'])}</b>\n"
        f"📊 Eski ovozlar: <b>{old_total}</b>\n"
        f"➕ Qo'shilayotgan ovozlar: <b>{sign}{amount}</b>\n"
        f"📊 Yangi jami: <b>{new_total}</b>\n\n"
        f"📝 Sabab: {escape(reason)}\n\n"
        "Tasdiqlaysizmi?",
        reply_markup=kb.adjust_confirm_keyboard(),
    )


@router.callback_query(F.data == "adjust_confirm")
async def cb_adjust_confirm(callback: CallbackQuery, state: FSMContext):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    data = await state.get_data()
    video_id = data.get("adjust_video_id")
    amount = data.get("adjust_amount")
    reason = data.get("adjust_reason")

    if video_id is None or amount is None or not reason:
        await callback.answer()
        return

    await db.add_vote_adjustment(video_id, amount, callback.from_user.id, reason)
    await state.update_data(adjust_video_id=None, adjust_amount=None, adjust_reason=None)

    video = await db.get_video(video_id)
    await callback.message.edit_text(
        f"✅ Tuzatish qo'llanildi.\n\n"
        f"🎬 {escape(video['title'])}\n"
        f"📊 Yangi jami ovozlar: <b>{video['votes_count']}</b>"
    )
    await callback.answer()


@router.callback_query(F.data == "adjust_cancel")
async def cb_adjust_cancel_flow(callback: CallbackQuery, state: FSMContext):
    await state.update_data(adjust_video_id=None, adjust_amount=None, adjust_reason=None)
    await callback.message.edit_text("❌ Bekor qilindi, hech narsa o'zgartirilmadi.")
    await callback.answer()


@router.message(F.text == "📜 Tuzatishlar tarixi")
async def adjustments_history(message: Message):
    if not await is_admin(message.from_user.id):
        return
    adjustments = await db.get_all_adjustments()
    if not adjustments:
        await message.answer("📜 Hozircha hech qanday ovoz tuzatishi qilinmagan.")
        return

    for adj in adjustments[:30]:
        status = "❌ BEKOR QILINGAN" if adj["is_cancelled"] else "✅ Faol"
        sign = "+" if adj["amount"] > 0 else ""
        date = adj["created_at"].split("T")[0]
        text = (
            f"{status}\n"
            f"🎬 {escape(adj['video_title'])}\n"
            f"⚖️ {sign}{adj['amount']} ovoz\n"
            f"👤 Admin ID: <code>{adj['admin_id']}</code>\n"
            f"📝 Sabab: {escape(adj['reason'])}\n"
            f"📅 {date}"
        )
        reply_markup = None
        if not adj["is_cancelled"]:
            reply_markup = kb.cancel_adjustment_keyboard(adj["id"])
        await message.answer(text, reply_markup=reply_markup)


@router.callback_query(F.data.startswith("cancel_adj:"))
async def cb_cancel_adjustment(callback: CallbackQuery):
    if not await is_admin(callback.from_user.id):
        await callback.answer()
        return
    adjustment_id = int(callback.data.split(":")[1])
    await db.cancel_adjustment(adjustment_id, callback.from_user.id)
    await callback.message.edit_text(
        callback.message.text + "\n\n🗑 BU TUZATISH BEKOR QILINDI.",
    )
    await callback.answer("Tuzatish bekor qilindi.")


# Eslatma: "🏆 Natijalar" tugmasi handlers/videos.py da barcha foydalanuvchilar
# (jumladan admin) uchun bitta joyda ishlaydi — bu yerda takrorlanmaydi.
